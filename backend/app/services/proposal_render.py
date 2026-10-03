"""
Flexible proposal template rendering + export (HTML, PDF, Excel, Image).

Template placeholders (Jinja2):
  company.name, company.address  (from settings defaults)
  proposal.number, proposal.title, proposal.notes, proposal.currency,
  proposal.status, proposal.created_at, proposal.valid_until
    customer.name, customer.company_name, customer.code, customer.area, customer.branch,
    customer.email, customer.phone, customer.address, customer.tax_id
  price_list.name
  items[]: sku, name, description, image_url, quantity, unit, unit_price, tax_rate,
           line_subtotal, line_tax, line_total
  totals.subtotal, totals.tax, totals.grand
  today
"""
from __future__ import annotations
import io
from html.parser import HTMLParser
from datetime import datetime, timezone
from decimal import Decimal
from jinja2 import Environment, BaseLoader, select_autoescape
from reportlab.lib.pagesizes import A4, letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import mm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from PIL import Image as PILImage, ImageDraw, ImageFont
from app.models.proposal import Proposal, ProposalTemplate

DEFAULT_HTML = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8"/>
<title>{{ proposal.number }}</title>
<style>
  {{ css }}
  body { font-family: Arial, sans-serif; color: #1e293b; margin: 24px; font-size: 12px; }
  h1 { font-size: 22px; margin: 0 0 4px; }
  .muted { color: #64748b; }
  .header { display: flex; justify-content: space-between; margin-bottom: 24px; }
  table { width: 100%; border-collapse: collapse; margin-top: 16px; }
  th, td { border-bottom: 1px solid #e2e8f0; padding: 8px; text-align: left; vertical-align: top; }
  th { background: #f8fafc; font-size: 11px; text-transform: uppercase; color: #64748b; }
  .right { text-align: right; }
  .totals { margin-top: 16px; width: 280px; margin-left: auto; }
  .totals td { border: none; padding: 4px 8px; }
  .grand { font-weight: bold; font-size: 14px; }
  img.prod { max-width: 64px; max-height: 64px; object-fit: contain; }
</style>
</head>
<body>
  <div class="header">
    <div>
      <h1>{{ company.name }}</h1>
      <div class="muted">{{ company.address }}</div>
    </div>
    <div class="right">
      <div style="font-size:18px;font-weight:bold;">PROPOSAL</div>
      <div>{{ proposal.number }}</div>
      <div class="muted">{{ proposal.created_at }}</div>
    </div>
  </div>
  <div style="margin-bottom:16px;">
    <strong>Customer:</strong> {{ customer.name }} ({{ customer.code }})<br/>
    {% if customer.company_name %}<strong>Company:</strong> {{ customer.company_name }}<br/>{% endif %}
    {% if customer.area %}<strong>Area:</strong> {{ customer.area }}<br/>{% endif %}
    {% if customer.branch %}<strong>Branch:</strong> {{ customer.branch }}<br/>{% endif %}
    {% if customer.address %}{{ customer.address }}<br/>{% endif %}
    {% if customer.email %}{{ customer.email }}{% endif %}
    {% if price_list %}<br/><span class="muted">Price list: {{ price_list.name }}</span>{% endif %}
  </div>
  {% if proposal.title %}<h2>{{ proposal.title }}</h2>{% endif %}
  <table>
    <thead>
      <tr>
        <th></th>
        <th>SKU</th>
        <th>Product</th>
        <th class="right">Qty</th>
        <th class="right">Unit</th>
        <th class="right">Tax%</th>
        <th class="right">Line total</th>
      </tr>
    </thead>
    <tbody>
    {% for it in items %}
      <tr>
        <td>{% if it.image_url %}<img class="prod" src="{{ it.image_url }}"/>{% endif %}</td>
        <td>{{ it.sku }}</td>
        <td>{{ it.name }}{% if it.description %}<div class="muted">{{ it.description }}</div>{% endif %}</td>
        <td class="right">{{ it.quantity }} {{ it.unit }}</td>
        <td class="right">{{ '%.2f'|format(it.unit_price) }}</td>
        <td class="right">{{ '%.2f'|format(it.tax_rate) }}</td>
        <td class="right">{{ '%.2f'|format(it.line_total) }}</td>
      </tr>
    {% endfor %}
    </tbody>
  </table>
  <table class="totals">
    <tr><td>Subtotal</td><td class="right">{{ '%.2f'|format(totals.subtotal) }} {{ proposal.currency }}</td></tr>
        <tr><td>Discount ({{ '%.2f'|format(totals.discount_percent) }}%)</td><td class="right">-{{ '%.2f'|format(totals.discount) }} {{ proposal.currency }}</td></tr>
    <tr><td>Tax</td><td class="right">{{ '%.2f'|format(totals.tax) }}</td></tr>
    <tr class="grand"><td>Total</td><td class="right">{{ '%.2f'|format(totals.grand) }} {{ proposal.currency }}</td></tr>
  </table>
  {% if proposal.notes %}<p class="muted" style="margin-top:24px;">{{ proposal.notes }}</p>{% endif %}
</body>
</html>
"""


def build_context(proposal: Proposal, company: dict | None = None) -> dict:
    items = []
    subtotal = 0.0
    tax_total = 0.0
    discount_percent = float(proposal.discount_percent or 0)
    for it in proposal.items:
        up = float(it.unit_price or 0)
        qty = float(it.quantity or 0)
        tr = float(it.tax_rate or 0)
        line_sub = up * qty
        line_discount = line_sub * discount_percent / 100.0
        discounted_subtotal = line_sub - line_discount
        line_tax = discounted_subtotal * tr / 100.0
        line_total = discounted_subtotal + line_tax
        subtotal += line_sub
        tax_total += line_tax
        items.append({
            "sku": it.sku or "",
            "name": it.name,
            "description": it.description or "",
            "image_url": it.image_url or (it.product.photo_url if it.product else "") or "",
            "quantity": qty,
            "unit": it.unit or "pcs",
            "unit_price": up,
            "tax_rate": tr,
            "line_subtotal": round(line_sub, 2),
            "line_discount": round(line_discount, 2),
            "line_tax": round(line_tax, 2),
            "line_total": round(line_total, 2),
        })
    discount_amount = subtotal * discount_percent / 100.0
    cust = proposal.customer
    return {
        "company": company or {"name": "Warehouse Co.", "address": ""},
        "proposal": {
            "number": proposal.number,
            "title": proposal.title or "",
            "notes": proposal.notes or "",
            "currency": proposal.currency or "TRY",
            "status": proposal.status.value if proposal.status else "",
            "created_at": proposal.created_at.strftime("%Y-%m-%d") if proposal.created_at else "",
            "valid_until": proposal.valid_until.strftime("%Y-%m-%d") if proposal.valid_until else "",
        },
        "customer": {
            "name": cust.name if cust else "",
            "company_name": cust.company_name or "" if cust else "",
            "code": cust.code if cust else "",
            "area": cust.area or "" if cust else "",
            "branch": cust.branch or "" if cust else "",
            "email": cust.email if cust else "",
            "phone": cust.phone if cust else "",
            "address": cust.address if cust else "",
            "tax_id": cust.tax_id if cust else "",
        },
        "price_list": {"name": proposal.price_list.name} if proposal.price_list else None,
        "items": items,
        "totals": {
            "subtotal": round(subtotal, 2),
            "discount": round(discount_amount, 2),
            "discount_percent": discount_percent,
            "tax": round(tax_total, 2),
            "grand": round(subtotal - discount_amount + tax_total, 2),
        },
        "today": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        "css": "",
    }


def render_html(proposal: Proposal, template: ProposalTemplate | None) -> str:
    ctx = build_context(proposal)
    body = (template.body_html if template and template.body_html else DEFAULT_HTML)
    css = (template.css if template and template.css else "")
    ctx["css"] = css or ""
    env = Environment(loader=BaseLoader(), autoescape=select_autoescape(["html", "xml"]))
    return env.from_string(body).render(**ctx)


def render_template_preview(body_html: str, css: str | None = None) -> str:
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    context = {
        "company": {"name": "Warehouse Co.", "address": "Industrial Zone 1"},
        "proposal": {
            "number": "PR-20261002-0001",
            "title": "Sample proposal",
            "notes": "Thank you for your business.",
            "currency": "TRY",
            "status": "draft",
            "created_at": today,
            "valid_until": "",
        },
        "customer": {
            "name": "Sample Customer",
            "company_name": "Sample Company Ltd.",
            "code": "CUST-001",
            "area": "Central",
            "branch": "Main branch",
            "email": "customer@example.com",
            "phone": "+90 555 010 0000",
            "address": "Example Street, Istanbul",
            "tax_id": "",
        },
        "price_list": {"name": "Standard"},
        "items": [{
            "sku": "SAMPLE-001",
            "name": "Sample product",
            "description": "Example line item",
            "image_url": "",
            "quantity": 2,
            "unit": "pcs",
            "unit_price": 50.0,
            "tax_rate": 20.0,
            "line_subtotal": 100.0,
            "line_tax": 20.0,
            "line_total": 120.0,
        }],
        "totals": {"subtotal": 100.0, "discount": 0.0, "discount_percent": 0.0, "tax": 20.0, "grand": 120.0},
        "today": today,
        "css": css or "",
    }
    env = Environment(loader=BaseLoader(), autoescape=select_autoescape(["html", "xml"]))
    return env.from_string(body_html or DEFAULT_HTML).render(**context)



def _html_to_text_blocks(html: str) -> list[str]:
    import re as _re
    html = _re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    html = _re.sub(r"(?is)<style.*?>.*?</style>", " ", html)
    html = _re.sub(r"(?i)<br\s*/?>", "\n", html)
    html = _re.sub(r"(?i)</p>", "\n", html)
    html = _re.sub(r"(?i)</tr>", "\n", html)
    html = _re.sub(r"(?i)</h[1-6]>", "\n", html)
    html = _re.sub(r"<[^>]+>", " ", html)
    html = _re.sub(r"[ \t]+", " ", html)
    return [ln.strip() for ln in html.split("\n") if ln.strip()]


def export_pdf(proposal: Proposal, template: ProposalTemplate | None) -> bytes:
    """Structured PDF; merges customer/items from context and template-rendered text."""
    ctx = build_context(proposal)
    html = render_html(proposal, template)
    buf = io.BytesIO()
    pagesize = A4
    if template and (template.page_size or "").lower() == "letter":
        pagesize = letter
    doc = SimpleDocTemplate(
        buf, pagesize=pagesize,
        leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("T", parent=styles["Heading1"], fontSize=16, spaceAfter=6)
    small = ParagraphStyle("S", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#64748b"))
    story = []
    story.append(Paragraph(f"PROPOSAL {ctx['proposal']['number']}", title))
    if ctx["proposal"]["title"]:
        story.append(Paragraph(str(ctx["proposal"]["title"]), styles["Heading2"]))
    story.append(Paragraph(
        f"<b>{ctx['customer']['name']}</b> ({ctx['customer']['code']})",
        styles["Normal"],
    ))
    if ctx.get("price_list"):
        story.append(Paragraph(f"Price list: {ctx['price_list']['name']}", small))
    story.append(Spacer(1, 8))
    data = [["SKU", "Product", "Qty", "Unit", "Tax%", "Total"]]
    for it in ctx["items"]:
        data.append([
            it["sku"],
            (it["name"] or "")[:48],
            f"{it['quantity']} {it['unit']}",
            f"{it['unit_price']:.2f}",
            f"{it['tax_rate']:.1f}",
            f"{it['line_total']:.2f}",
        ])
    data.append(["", "", "", "", "Subtotal", f"{ctx['totals']['subtotal']:.2f}"])
    data.append(["", "", "", "", f"Discount ({ctx['totals']['discount_percent']:.2f}%)", f"-{ctx['totals']['discount']:.2f}"])
    data.append(["", "", "", "", "Tax", f"{ctx['totals']['tax']:.2f}"])
    data.append(["", "", "", "", "Grand", f"{ctx['totals']['grand']:.2f} {ctx['proposal']['currency']}"])
    tbl = Table(data, colWidths=[65, 170, 55, 55, 40, 65])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -4), 0.25, colors.HexColor("#e2e8f0")),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("FONTNAME", (4, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    story.append(tbl)
    if ctx["proposal"]["notes"]:
        story.append(Spacer(1, 10))
        story.append(Paragraph(str(ctx["proposal"]["notes"]), styles["Normal"]))
    blocks = _html_to_text_blocks(html)[:8]
    if blocks:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Template content", small))
        for b in blocks[:6]:
            safe = b[:120].replace("&", "&amp;").replace("<", "&lt;")
            story.append(Paragraph(safe, small))
    doc.build(story)
    buf.seek(0)
    return buf.read()


class _PreviewTableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self.table_depth = 0
        self.table: list[list[str]] = []
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.headings: list[str] = []
        self.heading: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            if self.table_depth == 0:
                self.table = []
            self.table_depth += 1
        elif self.table_depth and tag == "tr":
            self.row = []
        elif self.table_depth and tag in ("td", "th"):
            self.cell = []
        elif self.table_depth and tag == "br" and self.cell is not None:
            self.cell.append(" ")
        elif not self.table_depth and tag in ("h1", "h2", "h3", "p"):
            self.heading = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)
        elif self.heading is not None:
            self.heading.append(data)

    def handle_endtag(self, tag):
        if self.table_depth and tag in ("td", "th") and self.cell is not None:
            value = " ".join("".join(self.cell).split())
            if self.row is not None:
                self.row.append(value)
            self.cell = None
        elif self.table_depth and tag == "tr" and self.row is not None:
            if self.row:
                self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table_depth:
            self.table_depth -= 1
            if self.table_depth == 0:
                self.tables.append(self.table)
        elif not self.table_depth and tag in ("h1", "h2", "h3", "p") and self.heading is not None:
            value = " ".join("".join(self.heading).split())
            if value:
                self.headings.append(value)
            self.heading = None


def export_excel(rendered_html: str) -> bytes:
    parser = _PreviewTableParser()
    parser.feed(rendered_html)
    wb = Workbook()
    ws = wb.active
    ws.title = "Proposal preview"
    for heading in parser.headings:
        ws.append([heading])
    for table_index, table in enumerate(parser.tables):
        if table_index or parser.headings:
            ws.append([])
        for row_index, row in enumerate(table):
            ws.append(row)
            is_header = row_index == 0 and any(value.strip().lower() in {"sku", "product", "qty", "line total"} for value in row)
            if is_header:
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="E2E8F0")
            elif any(value.lower().startswith("total") for value in row if isinstance(value, str)):
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True)
    for column in ws.columns:
        width = min(max(max((len(str(cell.value or "")) for cell in column), default=0) + 2, 10), 48)
        ws.column_dimensions[column[0].column_letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_image(proposal: Proposal, width: int = 900) -> bytes:
    """Render a simple summary image with Pillow (for sharing / preview)."""
    ctx = build_context(proposal)
    row_h = 28
    header_h = 140
    height = header_h + 40 + row_h * (len(ctx["items"]) + 4) + 80
    height = max(height, 400)
    img = PILImage.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 14)
        font_b = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 18)
        font_s = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 11)
    except Exception:
        font = font_b = font_s = ImageFont.load_default()

    draw.rectangle([0, 0, width, 64], fill=(15, 23, 42))
    draw.text((24, 18), f"PROPOSAL  {ctx['proposal']['number']}", fill=(255, 255, 255), font=font_b)
    y = 80
    draw.text((24, y), ctx["customer"]["name"], fill=(15, 23, 42), font=font_b)
    y += 28
    if ctx["proposal"]["title"]:
        draw.text((24, y), ctx["proposal"]["title"], fill=(71, 85, 105), font=font)
        y += 24
    y += 12
    draw.line([(24, y), (width - 24, y)], fill=(226, 232, 240), width=1)
    y += 12
    for it in ctx["items"]:
        line = f"{it['sku']}  {it['name'][:40]}  ×{it['quantity']}  {it['line_total']:.2f}"
        draw.text((24, y), line, fill=(30, 41, 59), font=font_s)
        y += row_h
    y += 8
    draw.line([(24, y), (width - 24, y)], fill=(226, 232, 240), width=1)
    y += 16
    draw.text((24, y), f"Subtotal: {ctx['totals']['subtotal']:.2f}   Discount: {ctx['totals']['discount']:.2f}   Tax: {ctx['totals']['tax']:.2f}", fill=(71, 85, 105), font=font)
    y += 28
    draw.text((24, y), f"TOTAL: {ctx['totals']['grand']:.2f} {ctx['proposal']['currency']}", fill=(15, 23, 42), font=font_b)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf.read()
